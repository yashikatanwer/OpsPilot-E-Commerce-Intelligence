import os
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd
from flask import Flask, render_template, request, redirect, url_for, session, send_file
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
import json

from utils.analysis import analyze_dataset, detect_fields, classify_ecommerce_dataset

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("OPSPILOT_SECRET_KEY", "opspilot-final-ecommerce")
USERS_FILE = os.path.join(BASE_DIR, "users.json")


def _load_users():
    if not os.path.exists(USERS_FILE):
        return {}
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as file:
            return json.load(file)
    except (OSError, json.JSONDecodeError):
        return {}


def _save_users(users):
    with open(USERS_FILE, "w", encoding="utf-8") as file:
        json.dump(users, file, indent=2)


def _require_login():
    if not session.get("logged_in"):
        return redirect(url_for("login"))
    return None

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024
ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".csv"}


def allowed_file(filename):
    return bool(filename) and os.path.splitext(filename)[1].lower() in ALLOWED_EXTENSIONS


def _normal_columns(columns):
    from utils.analysis import normalize_name
    return [normalize_name(c) for c in columns]


def _read_excel_sheets(filepath):
    """Read only compatible e-commerce sheets from a workbook.

    XLSX files use openpyxl's read-only iterator so large workbooks do not
    require the normal pandas Excel parser to materialize the entire workbook
    before analysis. Unrelated sheets are rejected at header level.
    """
    extension = os.path.splitext(filepath)[1].lower()
    if extension == ".csv":
        df = pd.read_csv(filepath)
        df.attrs["source_sheets"] = ["CSV"]
        return df

    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(filepath, read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f"Excel file could not be opened: {exc}") from exc

        candidates = []
        try:
            # First pass: inspect headers only.
            for ws in wb.worksheets:
                iterator = ws.iter_rows(values_only=True)
                try:
                    header = next(iterator)
                except StopIteration:
                    continue
                headers = [str(x).strip() if x not in (None, "") else f"Unnamed: {i}" for i, x in enumerate(header)]
                header_df = pd.DataFrame(columns=headers)
                fields = detect_fields(header_df)
                classification = classify_ecommerce_dataset(header_df, fields)
                if classification["is_ecommerce"]:
                    candidates.append((ws.title, headers, fields))

            if not candidates:
                raise ValueError("No supported e-commerce order/sales sheet was found in this workbook.")

            first_name, first_headers, first_fields = candidates[0]
            first_raw = set(_normal_columns(first_headers))
            first_semantic = {k for k, v in first_fields.items() if v}
            selected = []
            for name, headers, fields in candidates:
                raw_overlap = len(first_raw & set(_normal_columns(headers)))
                semantic_overlap = len(first_semantic & {k for k, v in fields.items() if v})
                if name == first_name or semantic_overlap >= 3 or raw_overlap >= 3:
                    selected.append((name, headers))

            frames = []
            for sheet_name, headers in selected:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or all(v is None for v in row):
                        continue
                    values = list(row[:len(headers)])
                    if len(values) < len(headers):
                        values.extend([None] * (len(headers) - len(values)))
                    rows.append(values)
                if rows:
                    frames.append(pd.DataFrame(rows, columns=headers))

            if not frames:
                raise ValueError("The supported e-commerce sheet(s) contain no data rows.")
            df = pd.concat(frames, ignore_index=True, sort=False)
            df = df.dropna(axis=1, how="all").dropna(axis=0, how="all")
            df.attrs["source_sheets"] = [name for name, _ in selected]
            return df
        finally:
            wb.close()

    # Legacy .xls support through xlrd/pandas.
    try:
        book = pd.ExcelFile(filepath)
        usable = []
        for sheet_name in book.sheet_names:
            sheet = pd.read_excel(book, sheet_name=sheet_name)
            if sheet is None or sheet.empty:
                continue
            sheet = sheet.dropna(axis=1, how="all").dropna(axis=0, how="all")
            fields = detect_fields(sheet)
            if classify_ecommerce_dataset(sheet, fields)["is_ecommerce"]:
                usable.append((str(sheet_name), sheet, fields))
        book.close()
    except Exception as exc:
        raise ValueError(f".xls file could not be read. Install xlrd with the project requirements. Details: {exc}") from exc

    if not usable:
        raise ValueError("No supported e-commerce order/sales sheet was found in this workbook.")
    first_name, first_df, first_fields = usable[0]
    first_raw = set(_normal_columns(first_df.columns))
    first_semantic = {k for k, v in first_fields.items() if v}
    selected = [first_df]
    names = [first_name]
    for name, sheet, fields in usable[1:]:
        if len(first_raw & set(_normal_columns(sheet.columns))) >= 3 or len(first_semantic & {k for k, v in fields.items() if v}) >= 3:
            selected.append(sheet)
            names.append(name)
    df = pd.concat(selected, ignore_index=True, sort=False)
    df.attrs["source_sheets"] = names
    return df


def load_dataset(filepath):
    df = _read_excel_sheets(filepath)
    if df.empty:
        raise ValueError("The uploaded dataset is empty.")
    return df


def _current_file():
    if not session.get("dataset_uploaded"):
        return None
    filename = session.get("dataset_filename")
    if filename:
        path = os.path.join(UPLOAD_FOLDER, filename)
        if os.path.isfile(path):
            return path
    # Recovery if the session has a stale filename.
    files = [os.path.join(UPLOAD_FOLDER, f) for f in os.listdir(UPLOAD_FOLDER) if allowed_file(f)]
    return max(files, key=os.path.getmtime) if files else None


def get_analysis():
    filepath = _current_file()
    if not filepath:
        session.pop("dataset_uploaded", None)
        session.pop("dataset_filename", None)
        return None, None, None
    df = load_dataset(filepath)
    df.attrs["filename"] = os.path.basename(filepath)
    summary = analyze_dataset(df)
    summary["filename"] = os.path.basename(filepath)
    summary["source_sheets"] = df.attrs.get("source_sheets", [])
    return filepath, df, summary


def _page_context(filepath, df, summary):
    return {
        "filename": os.path.basename(filepath) if filepath else None,
        "data": df,
        "summary": summary,
        "user_name": session.get("user_name", "User"),
        "user_email": session.get("user_email", "")
    }


@app.before_request
def require_login():
    public_endpoints = {"login", "signup", "static"}
    if request.endpoint not in public_endpoints and not session.get("logged_in"):
        return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("logged_in"):
        return redirect(url_for("dashboard"))

    error = None
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        users = _load_users()
        user = users.get(email)

        if user and check_password_hash(user["password"], password):
            session["logged_in"] = True
            session["user_email"] = email
            session["user_name"] = user.get("name", "User")
            return redirect(url_for("dashboard"))
        error = "Incorrect email or password."

    return render_template("login.html", error=error)


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if session.get("logged_in"):
        return redirect(url_for("dashboard"))

    error = None
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not name or not email or not password:
            error = "Please fill in all fields."
        elif "@" not in email or "." not in email.split("@")[-1]:
            error = "Please enter a valid email address."
        elif len(password) < 6:
            error = "Password must be at least 6 characters."
        elif password != confirm_password:
            error = "Passwords do not match."
        else:
            users = _load_users()
            if email in users:
                error = "An account with this email already exists."
            else:
                users[email] = {
                    "name": name,
                    "password": generate_password_hash(password)
                }
                _save_users(users)
                session["logged_in"] = True
                session["user_email"] = email
                session["user_name"] = name
                return redirect(url_for("dashboard"))

    return render_template("signup.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
def home():
    return redirect(url_for("dashboard"))


@app.route("/dashboard")
def dashboard():
    filepath, df, summary = get_analysis()
    return render_template("dashboard.html", **_page_context(filepath, df, summary))


@app.route("/upload", methods=["GET", "POST"])
def upload():
    if request.method == "GET":
        return render_template("upload.html")

    file = request.files.get("file")
    if not file or not file.filename:
        return render_template("upload.html", error="Please choose an Excel or CSV file.")
    if not allowed_file(file.filename):
        return render_template("upload.html", error="Unsupported file type. Use .xlsx, .xls or .csv.")

    # Replace the previous dataset only after the new file has been successfully saved/analyzed.
    filename = secure_filename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    temp_path = filepath + ".tmp"

    try:
        file.save(temp_path)
        df = load_dataset(temp_path)
        df.attrs["filename"] = filename
        summary = analyze_dataset(df)
        os.replace(temp_path, filepath)

        for old in os.listdir(UPLOAD_FOLDER):
            old_path = os.path.join(UPLOAD_FOLDER, old)
            if old_path != filepath and os.path.isfile(old_path) and allowed_file(old):
                try:
                    os.remove(old_path)
                except OSError:
                    pass

        session["dataset_uploaded"] = True
        session["dataset_filename"] = filename
        summary["filename"] = filename
        summary["source_sheets"] = df.attrs.get("source_sheets", [])

        table = df.head(10).to_html(classes="data-table", index=False, border=0, na_rep="—")
        return render_template("table.html", summary=summary, filename=filename, table=table)
    except Exception as exc:
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass
        return render_template("upload.html", error=str(exc))


@app.route("/analytics")
def analytics():
    filepath, df, summary = get_analysis()
    if summary is None:
        return redirect(url_for("upload"))
    return render_template("analytics.html", **_page_context(filepath, df, summary))


@app.route("/ai-insights")
def ai_insights():
    filepath, df, summary = get_analysis()
    if summary is None:
        return redirect(url_for("upload"))
    return render_template("insights.html", **_page_context(filepath, df, summary))


@app.route("/reports")
def reports():
    filepath, df, summary = get_analysis()
    if summary is None:
        return redirect(url_for("upload"))
    return render_template("reports.html", **_page_context(filepath, df, summary))


@app.route("/download-report")
def download_report():
    filepath, df, summary = get_analysis()
    if summary is None:
        return redirect(url_for("upload"))
    if not REPORTLAB_AVAILABLE:
        return render_template("reports.html", **_page_context(filepath, df, summary), error="PDF export requires the reportlab package. Run pip install -r requirements.txt.")

    kpis = summary.get("kpis", {})
    buffer = BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = [
        Paragraph("OpsPilot — E-Commerce Business Intelligence Report", styles["Title"]),
        Spacer(1, 10),
        Paragraph(f"Dataset: {escape(os.path.basename(filepath))}", styles["BodyText"]),
        Paragraph(f"Rows analyzed: {summary['rows']:,} | Columns: {summary['columns']}", styles["BodyText"]),
        Spacer(1, 16),
        Paragraph("Key Performance Indicators", styles["Heading2"]),
    ]

    rows = [["Metric", "Value"]]
    labels = [
        ("total_revenue", "Net revenue"), ("gross_revenue", "Gross revenue"), ("return_value", "Returns / refunds"),
        ("total_orders", "Orders"), ("total_customers", "Customers"), ("average_order_value", "Average order value"),
        ("total_quantity", "Quantity"), ("total_profit", "Profit"), ("cancelled_orders", "Cancelled / returned orders"),
        ("cancellation_rate", "Cancellation / return rate"),
    ]
    for key, label in labels:
        if key in kpis and kpis[key] is not None:
            value = kpis[key]
            if key in {"total_revenue", "gross_revenue", "return_value", "average_order_value", "total_profit"}:
                value = f"₹{value:,.2f}"
            elif key == "cancellation_rate":
                value = f"{value:.2f}%"
            else:
                value = f"{value:,}" if isinstance(value, (int, float)) else str(value)
            rows.append([label, value])
    table = Table(rows, colWidths=[250, 220])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b4f88")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c8d8e8")),
        ("PADDING", (0, 0), (-1, -1), 7),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story += [table, Spacer(1, 18), Paragraph("Detected structure", styles["Heading2"])]
    for field, column in summary.get("detected_fields", {}).items():
        if column:
            story.append(Paragraph(f"<b>{escape(field.replace('_', ' ').title())}:</b> {escape(str(column))}", styles["BodyText"]))
    story += [Spacer(1, 14), Paragraph("Business insights", styles["Heading2"])]
    for item in summary.get("insights", []):
        story.append(Paragraph(f"<b>{escape(str(item['title']))}:</b> {escape(str(item['text']))}", styles["BodyText"]))
        story.append(Spacer(1, 5))

    doc.build(story)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name="OpsPilot_Ecommerce_Report.pdf", mimetype="application/pdf")


@app.route("/reset")
def reset():
    for filename in os.listdir(UPLOAD_FOLDER):
        path = os.path.join(UPLOAD_FOLDER, filename)
        if os.path.isfile(path):
            try:
                os.remove(path)
            except OSError:
                pass
    session.clear()
    return redirect(url_for("upload"))


if __name__ == "__main__":
    app.run(debug=True)
